"""Generator syntetycznych fikstur wiernych formatom z §2 specyfikacji.

Brak realnych plików wejściowych w repo, więc budujemy spójny, deterministyczny
zestaw danych dla okresu 2026-02, który odwzorowuje WSZYSTKIE golden case'y z §9:

* dubel wewnątrz-przewoźnika SPT ``Shoper46046-1`` (259,60 + 583,00 = 842,60),
* dubel cross-carrier ``47559`` (Zadbano 568,52 + D&M 455,00),
* 12 zleceń „niepowodzenie" u Zadbano (Σ 5 251,37) + 3 „anulowane",
* złączenia MYBED43503->Shoper43503-1, Shoper46281-1_8801122->Shoper46281-1,
  47559->Shoper47559-1,
* zdublowana treść PDF D&M (dedup),
* twarde sumy uzgodnienia: SPT 29 225,96 / Zadbano 98 172,50 / D&M 5 445,00.

Pliki są w pełni deterministyczne (bez RNG), więc nadają się do commitu.
Uruchom: ``python -m transport_audit.tests.fixtures.generate``.
"""

from __future__ import annotations

import csv
import json
import shutil
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

FIX_DIR = Path(__file__).resolve().parent

# Rejestracja fontu z obsługą polskich znaków (Helvetica ich nie ma).
_FONT = "Helvetica"
_FONT_BOLD = "Helvetica-Bold"
for _cand in ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",):
    if Path(_cand).exists():
        try:
            pdfmetrics.registerFont(TTFont("PLSans", _cand))
            _bold = _cand.replace("DejaVuSans.ttf", "DejaVuSans-Bold.ttf")
            if Path(_bold).exists():
                pdfmetrics.registerFont(TTFont("PLSans-Bold", _bold))
            else:
                pdfmetrics.registerFont(TTFont("PLSans-Bold", _cand))
            _FONT, _FONT_BOLD = "PLSans", "PLSans-Bold"
        except Exception:
            pass
        break

PERIOD = "2026-02"
SPT_INVOICE = "FS/38/02/2026"
ZADBANO_INVOICE = "307/02/2026/TR"
ZADBANO_SETTLEMENT = "307/02/2026/TR"
DM_INVOICE = "FS/24/02/2026"
DM_SETTLEMENT = "07/02/2026"

SPT_TARGET = 29225.96
ZADBANO_TARGET = 98172.50
DM_TARGET = 5445.00

# ERP kolumny (§2.1) — dokładna kolejność 24 kolumn
ERP_COLUMNS = [
    "Numer", "Data zamówienia", "Przewidywana data",
    "Pozycje zamówienia/Produkt/Nazwa", "Pozycje zamówienia/Ilość",
    "Pozycje zamówienia/Opcja", "Metoda dostawy", "Kod pocztowy", "Miasto",
    "Suma", "Koszt dostawy", "Uwagi", "Zapłacone", "Status", "Tagi",
    "Data realizacji", "Kupon rabatowy", "Faktura produkcyjna",
    "Faktura materac", "Faktura transportowa", "Klient/Nazwa", "Klient/E-mail",
    "Klient/Telefon", "Adres dostawy/Ulica",
]

SERVICE_ITEM = {
    "DOOR": "Wysyłka",
    "CARRY_IN": "Wniesienie zamówienia",
    "CARRY_IN_ASSEMBLY": "Wniesienie zamówienia z usługą montażu i sprzątania",
}


def pl(amount: float) -> str:
    """Sformatuj kwotę po polsku: 1234.5 -> '1 234,50'."""
    s = f"{amount:,.2f}"
    return s.replace(",", " ").replace(".", ",")


# --------------------------------------------------------------------------- #
# Model scenariusza
# --------------------------------------------------------------------------- #


@dataclass
class ErpOrderSpec:
    number: str
    receiver: str
    postcode: str
    city: str
    street: str
    order_date: str = "2026-02-05"
    service_level: str = "DOOR"
    products: list[tuple[str, int, str]] = field(default_factory=list)  # (name, qty, opcja)
    delivery_fee: float = 149.0
    order_total: float = 2499.0
    status: str = "zrealizowane"
    tags: str = ""
    transport_invoice: str = ""   # istniejący ręczny wpis (backtest §9)
    delivery_method: str = "SPT Logistics"
    email: str = "klient@example.com"
    phone: str = "600100200"


ERP: dict[str, ErpOrderSpec] = {}


def erp_order(number: str, receiver: str, postcode: str, city: str, **kw) -> ErpOrderSpec:
    """Zarejestruj (lub zwróć) zamówienie ERP; domyślnie 1 łóżko 160x200."""
    if number in ERP:
        return ERP[number]
    street = kw.pop("street", "ul. Testowa 1")
    spec = ErpOrderSpec(number=number, receiver=receiver, postcode=postcode,
                        city=city, street=street, **kw)
    if not spec.products:
        spec.products = [("Łóżko kontynentalne Lea", 1, "Powierzchnia spania: 160cm x 200cm")]
    ERP[number] = spec
    return spec


@dataclass
class Comp:
    type_raw: str
    amount: float


@dataclass
class ZLine:
    order_ref: str
    status: str          # zrealizowane / niepowodzenie / anulowane / w realizacji
    standard: str        # Dostawa / Dostawa Standard / Dostawa Komfort
    receiver: str
    city: str
    postcode: str
    comps: list[Comp]
    qty: float = 1
    weight: float = 40.0
    volume: float = 0.8
    waybill: str = ""


@dataclass
class FlatLine:
    """Wspólny wiersz „flat" dla SPT / D&M."""
    order_ref: str
    receiver: str
    postcode: str
    city: str
    products: str
    transport_net: float
    status_label: str          # SPT: Odebrane/Nieodebrane ; D&M: (n/d)
    order_type: str = "Zam. zwykłe"   # SPT `Typ zamówienia`
    services_net: float = 0.0         # D&M `Usługi netto`
    weight: float = 40.0
    volume: float = 0.8
    packages: int = 2


# --------------------------------------------------------------------------- #
# ZADBANO — zestawienie itemized (target 98 172,50)
# --------------------------------------------------------------------------- #

def build_zadbano() -> tuple[list[ZLine], dict]:
    lines: list[ZLine] = []

    # --- golden cross-carrier 47559 (Dominika Bronner) = 568,52 --------------
    erp_order("Shoper47559-1", "Dominika Bronner", "30-001", "Kraków",
              service_level="DOOR", transport_invoice="307/02/2026/TR ZADBANO",
              products=[("Łóżko Auro Twin", 1, "Powierzchnia spania: 90cm x 200cm")])
    lines.append(ZLine(
        order_ref="Shoper47559-1_8801122", status="zrealizowane", standard="Dostawa",
        receiver="Dominika Bronner", city="Kraków", postcode="30-001",
        comps=[Comp("Transport", 480.00), Comp("Dopłata paliwowa", 48.00),
               Comp("Opłata drogowa", 25.00), Comp("Dopłata sezonowa", 15.52)],
        weight=45, volume=0.9, waybill="ZAD-47559"))

    # --- 46281 komfort (wniesienie + montaż) --------------------------------
    erp_order("Shoper46281-1", "Jan Kowalski", "80-001", "Gdańsk",
              service_level="CARRY_IN_ASSEMBLY",
              transport_invoice="307/02/2026/TR ZADBANO")
    lines.append(ZLine(
        order_ref="Shoper46281-1_8801122", status="zrealizowane",
        standard="Dostawa Komfort", receiver="Jan Kowalski", city="Gdańsk",
        postcode="80-001",
        comps=[Comp("Transport", 400.00), Comp("Dopłata za standard dostawy", 120.00),
               Comp("Dopłata paliwowa", 40.00)], weight=90, volume=1.6,
        waybill="ZAD-46281"))

    # --- service mismatch: przewoźnik CARRY_IN, ERP DOOR --------------------
    erp_order("Shoper46600-1", "Anna Nowak", "50-001", "Wrocław",
              service_level="DOOR")
    lines.append(ZLine(
        order_ref="Shoper46600-1", status="zrealizowane", standard="Dostawa Komfort",
        receiver="Anna Nowak", city="Wrocław", postcode="50-001",
        comps=[Comp("Transport", 300.00), Comp("Dopłata za standard dostawy", 120.00)],
        weight=60, volume=1.0))

    # --- 12 „niepowodzenie" (Σ 5 251,37) ------------------------------------
    fail_values = [437.61] * 11 + [437.66]
    for i, val in enumerate(fail_values):
        num = f"Shoper4800{i:01d}-1" if i < 10 else f"Shoper480{i:02d}-1"
        num = f"Shoper{48001 + i}-1"
        erp_order(num, f"Klient Fail {i+1}", "02-100", "Warszawa")
        lines.append(ZLine(
            order_ref=num, status="niepowodzenie", standard="Dostawa",
            receiver=f"Klient Fail {i+1}", city="Warszawa", postcode="02-100",
            comps=[Comp("Transport", val)], weight=35, volume=0.7))

    # --- 3 „anulowane" obciążone --------------------------------------------
    for i, val in enumerate([300.00, 250.00, 200.00]):
        num = f"Shoper{48013 + i}-1"
        erp_order(num, f"Klient Anulo {i+1}", "61-001", "Poznań")
        lines.append(ZLine(
            order_ref=num, status="anulowane", standard="Dostawa",
            receiver=f"Klient Anulo {i+1}", city="Poznań", postcode="61-001",
            comps=[Comp("Transport", val)], weight=30, volume=0.6))

    # --- komponenty specjalne: Rabat / weryfikacja objętości / próba / zmiany
    erp_order("Shoper48020-1", "Klient Rabat", "31-001", "Kraków")
    lines.append(ZLine(
        order_ref="Shoper48020-1", status="zrealizowane", standard="Dostawa",
        receiver="Klient Rabat", city="Kraków", postcode="31-001",
        comps=[Comp("Transport", 500.00), Comp("Rabat", -50.00)], weight=50, volume=0.9))

    erp_order("Shoper48021-1", "Klient Objetosc", "80-010", "Gdańsk")
    lines.append(ZLine(
        order_ref="Shoper48021-1", status="zrealizowane", standard="Dostawa",
        receiver="Klient Objetosc", city="Gdańsk", postcode="80-010",
        comps=[Comp("Transport", 450.00), Comp("weryfikacja objętości", 60.00)],
        weight=55, volume=2.4))  # objętość mocno > ERP (0.8) -> volume upcharge

    erp_order("Shoper48022-1", "Klient Proba", "40-001", "Katowice")
    lines.append(ZLine(
        order_ref="Shoper48022-1", status="zrealizowane", standard="Dostawa",
        receiver="Klient Proba", city="Katowice", postcode="40-001",
        comps=[Comp("Transport", 420.00), Comp("dodatkowa próba doręczenia", 80.00),
               Comp("Zmiana daty - Vendor", 30.00), Comp("Zmiana adresu", 25.00)],
        weight=48, volume=0.85))

    # --- klaster cennikowy DOOR / postcode 90 (>=5 obs) + 1 przepłata -------
    cluster_vals = [450.00, 455.00, 448.00, 452.00, 460.00, 449.00]
    for i, val in enumerate(cluster_vals):
        num = f"Shoper{48100 + i}-1"
        erp_order(num, f"Klient Klaster {i+1}", "90-001", "Łódź")
        lines.append(ZLine(
            order_ref=num, status="zrealizowane", standard="Dostawa",
            receiver=f"Klient Klaster {i+1}", city="Łódź", postcode="90-001",
            comps=[Comp("Transport", val)], weight=40, volume=0.8))
    # przepłata: 600 > mediana(~451) * 1.20 = 541 -> FLAG above tariff
    erp_order("Shoper48120-1", "Klient Przeplata", "90-002", "Łódź")
    lines.append(ZLine(
        order_ref="Shoper48120-1", status="zrealizowane", standard="Dostawa",
        receiver="Klient Przeplata", city="Łódź", postcode="90-001",
        comps=[Comp("Transport", 600.00)], weight=40, volume=0.8))

    # --- suma dotychczasowa -------------------------------------------------
    def line_total(ln: ZLine) -> float:
        return round(sum(c.amount for c in ln.comps), 2)

    running = round(sum(line_total(l) for l in lines), 2)
    residual = round(ZADBANO_TARGET - running, 2)

    # --- filler: normalne DOOR zrealizowane, żeby trafić w 98 172,50 --------
    n_fill = 90
    base = round(residual / n_fill, 2)
    fill_vals = [base] * (n_fill - 1)
    fill_vals.append(round(residual - sum(fill_vals), 2))
    for i, val in enumerate(fill_vals):
        num = f"Shoper{49000 + i}-1"
        pc = f"{20 + (i % 60):02d}-{100 + i:03d}"
        erp_order(num, f"Klient Zad {i+1}", pc[:6], "Warszawa",
                  transport_invoice="307/02/2026/TR ZADBANO")
        lines.append(ZLine(
            order_ref=num, status="zrealizowane", standard="Dostawa",
            receiver=f"Klient Zad {i+1}", city="Warszawa", postcode=pc[:6],
            comps=[Comp("Transport", val)], weight=40, volume=0.8))

    total = round(sum(line_total(l) for l in lines), 2)
    assert abs(total - ZADBANO_TARGET) < 0.01, f"Zadbano total {total} != {ZADBANO_TARGET}"

    # podsumowanie per kategoria
    cats: dict[str, float] = {}
    for ln in lines:
        for c in ln.comps:
            cats[c.type_raw] = round(cats.get(c.type_raw, 0.0) + c.amount, 2)
    summary = {"categories": cats, "total": ZADBANO_TARGET}
    return lines, summary


# --------------------------------------------------------------------------- #
# D&M — zestawienie flat + usługi (target 5 445,00), z dedupem stron
# --------------------------------------------------------------------------- #

def build_dm() -> list[FlatLine]:
    lines: list[FlatLine] = []
    # golden cross-carrier 47559 = 455,00
    erp_order("Shoper47559-1", "Dominika Bronner", "30-001", "Kraków")  # już jest
    lines.append(FlatLine(order_ref="47559", receiver="Dominika Bronner",
                          postcode="30-001", city="Kraków",
                          products="Łóżko Auro Twin", transport_net=455.00,
                          status_label="", weight=45, volume=0.9, packages=2))

    dm_vals = [520.00, 480.00, 610.00, 390.00, 455.00, 700.00, 365.00, 430.00,
               640.00, 400.00]
    nums = [47560, 47561, 47562, 47563, 47564, 47565, 47566, 47567, 47568, 47569]
    products_pool = ["Łóżko kontynentalne + topper", "Łóżko Silla 2",
                     "Materac Hybrid 160", "Łóżko Auro King", "Topper 180"]
    for i, (num, val) in enumerate(zip(nums, dm_vals)):
        erp_order(f"Shoper{num}-1", f"Klient DM {i+1}", f"{60+i:02d}-100", "Poznań")
        lines.append(FlatLine(
            order_ref=str(num), receiver=f"Klient DM {i+1}",
            postcode=f"{60+i:02d}-100", city="Poznań",
            products=products_pool[i % len(products_pool)],
            transport_net=val, status_label="", weight=38 + i, volume=0.7 + i * 0.05,
            packages=2))

    total = round(sum(l.transport_net for l in lines), 2)
    assert abs(total - DM_TARGET) < 0.01, f"DM total {total} != {DM_TARGET}"
    return lines


# --------------------------------------------------------------------------- #
# SPT — zestawienie flat (target 29 225,96), z golden dublem 46046
# --------------------------------------------------------------------------- #

def build_spt() -> list[FlatLine]:
    lines: list[FlatLine] = []

    # golden dubel wewnątrz-przewoźnika: Shoper46046-1 (Nieodebrane 259,60 + Odebrane 583,00)
    erp_order("Shoper46046-1", "Piotr Zieliński", "70-001", "Szczecin",
              service_level="DOOR", transport_invoice="FS/38/02/2026 SPT")
    lines.append(FlatLine(order_ref="Shoper46046-1", receiver="Piotr Zieliński",
                          postcode="70-001", city="Szczecin",
                          products="1 szt. ERP_KONTYNENT_160", transport_net=259.60,
                          status_label="Nieodebrane", order_type="Zam. zwykłe"))
    lines.append(FlatLine(order_ref="Shoper46046-1", receiver="Piotr Zieliński",
                          postcode="70-001", city="Szczecin",
                          products="1 szt. ERP_KONTYNENT_160", transport_net=583.00,
                          status_label="Odebrane", order_type="Zam. zwykłe"))

    # złączenie MYBED43503 -> Shoper43503-1
    erp_order("Shoper43503-1", "Maria Wiśniewska", "00-001", "Warszawa",
              service_level="DOOR", transport_invoice="FS/38/02/2026 SPT")
    lines.append(FlatLine(order_ref="MYBED43503", receiver="Maria Wiśniewska",
                          postcode="00-001", city="Warszawa",
                          products="1 szt. ERP_KONTYNENT_160", transport_net=149.00,
                          status_label="Odebrane", order_type="Zam. zwykłe"))

    # złączenie wprost Shoper43183-1 (z wniesieniem -> CARRY_IN)
    erp_order("Shoper43183-1", "Tomasz Lewandowski", "31-500", "Kraków",
              service_level="CARRY_IN", transport_invoice="FS/38/02/2026 SPT")
    lines.append(FlatLine(order_ref="Shoper43183-1", receiver="Tomasz Lewandowski",
                          postcode="31-500", city="Kraków",
                          products="1 szt. ERP_KONTYNENT_160", transport_net=249.00,
                          status_label="Odebrane", order_type="Zam. z wniesieniem"))

    # service mismatch SPT: przewoźnik z wniesieniem, ERP DOOR
    erp_order("Shoper43900-1", "Ewa Dąbrowska", "20-001", "Lublin",
              service_level="DOOR")
    lines.append(FlatLine(order_ref="Shoper43900-1", receiver="Ewa Dąbrowska",
                          postcode="20-001", city="Lublin",
                          products="1 szt. ERP_KONTYNENT_180", transport_net=299.00,
                          status_label="Odebrane", order_type="Zam. z wniesieniem"))

    # kolizja prefiksów: core 12345 w ERP jako Shoper12345-1 i Shopify12345-1
    erp_order("Shoper12345-1", "Kolizja Shoper", "00-950", "Warszawa")
    erp_order("Shopify12345-1", "Kolizja Shopify", "80-950", "Gdańsk")
    lines.append(FlatLine(order_ref="Shoper12345-1", receiver="Kolizja Shoper",
                          postcode="00-950", city="Warszawa",
                          products="1 szt. ERP_KONTYNENT_160", transport_net=149.00,
                          status_label="Odebrane"))

    # core-match z prefiksem ZAM: ref 'ZAM/01847' -> core 01847 -> ERP 'ZAM01847'
    erp_order("ZAM01847", "Grzegorz Bąk", "44-100", "Gliwice", order_date="2026-02-10")
    lines.append(FlatLine(order_ref="ZAM/01847 - ODBIOR WLASNY", receiver="Grzegorz Bąk",
                          postcode="44-100", city="Gliwice",
                          products="1 szt. ERP_KONTYNENT_140", transport_net=180.00,
                          status_label="Odebrane"))

    # fuzzy fallback: brak numeru (core=None) -> dopasowanie po nazwisku + kodzie
    erp_order("Shopify77777-1", "Halina Górska", "43-300", "Bielsko-Biała",
              order_date="2026-02-12")
    lines.append(FlatLine(order_ref="ODBIOR WLASNY BIELSKO", receiver="Halina Górska",
                          postcode="43-300", city="Bielsko-Biała",
                          products="1 szt. ERP_KONTYNENT_160", transport_net=175.00,
                          status_label="Odebrane"))

    # orphan: linia bez odpowiednika w ERP
    lines.append(FlatLine(order_ref="MYBED ZWROT ID: 600794", receiver="Nieznany Klient",
                          postcode="99-999", city="Nigdzie",
                          products="zwrot", transport_net=120.00, status_label="Odebrane"))

    running = round(sum(l.transport_net for l in lines), 2)
    residual = round(SPT_TARGET - running, 2)
    n_fill = 60
    base = round(residual / n_fill, 2)
    fill_vals = [base] * (n_fill - 1)
    fill_vals.append(round(residual - sum(fill_vals), 2))
    for i, val in enumerate(fill_vals):
        num = f"Shoper{45000 + i}-1"
        pc = f"{10 + (i % 80):02d}-{200 + i:03d}"
        erp_order(num, f"Klient SPT {i+1}", pc[:6], "Kraków",
                  transport_invoice="FS/38/02/2026 SPT")
        lines.append(FlatLine(order_ref=num, receiver=f"Klient SPT {i+1}",
                              postcode=pc[:6], city="Kraków",
                              products="1 szt. ERP_KONTYNENT_160", transport_net=val,
                              status_label="Odebrane"))

    total = round(sum(l.transport_net for l in lines), 2)
    assert abs(total - SPT_TARGET) < 0.01, f"SPT total {total} != {SPT_TARGET}"
    return lines


# --------------------------------------------------------------------------- #
# Dodatkowe zamówienia ERP: unbilled + poza okresem
# --------------------------------------------------------------------------- #

def add_extra_erp() -> None:
    # unbilled: w okresie, ma Wysyłkę, brak linii kosztu u przewoźników
    erp_order("Shoper49999-1", "Klient Unbilled", "35-001", "Rzeszów",
              service_level="CARRY_IN", order_date="2026-02-15")
    # poza okresem (styczeń) — nie powinno trafiać do unbilled 2026-02
    erp_order("Shoper40000-1", "Klient Styczen", "00-002", "Warszawa",
              service_level="DOOR", order_date="2026-01-20")


# --------------------------------------------------------------------------- #
# Rendery plików
# --------------------------------------------------------------------------- #

def write_erp_csv(path: Path) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(ERP_COLUMNS)
        for spec in ERP.values():
            rows = _erp_rows(spec)
            for r in rows:
                writer.writerow(r)


def _erp_rows(spec: ErpOrderSpec) -> list[list[str]]:
    """Wygeneruj wiersze „long" dla jednego zamówienia (nagłówek + pozycje)."""
    # lista pozycji: produkty + pozycja-usługa wg poziomu
    items = list(spec.products)
    service_name = SERVICE_ITEM[spec.service_level]
    items.append((service_name, 1, ""))

    rows: list[list[str]] = []
    for idx, (name, qty, opcja) in enumerate(items):
        row = [""] * len(ERP_COLUMNS)
        if idx == 0:
            row[0] = spec.number
            row[1] = spec.order_date
            row[2] = spec.order_date
            row[6] = spec.delivery_method
            row[7] = spec.postcode
            row[8] = spec.city
            row[9] = pl(spec.order_total)
            row[10] = pl(spec.delivery_fee)
            row[11] = ""                       # Uwagi
            row[12] = pl(spec.order_total)     # Zapłacone
            row[13] = spec.status
            row[14] = spec.tags
            row[15] = spec.order_date          # Data realizacji
            row[16] = ""                       # Kupon
            row[17] = "FP/1/2026"              # Faktura produkcyjna
            row[18] = ""                       # Faktura materac
            row[19] = spec.transport_invoice   # Faktura transportowa (ręczna)
            row[20] = spec.receiver
            row[21] = spec.email
            row[22] = spec.phone
            row[23] = spec.street
        # kolumny pozycji zawsze
        row[3] = name
        row[4] = str(qty)
        row[5] = opcja
        rows.append(row)
    return rows


# ---- ZADBANO XLSX ---------------------------------------------------------- #

ZADBANO_HEADERS = [
    "LP", "Nr listu przewozowego", "Nr zamówienia", "Typ zlecenia",
    "Status realizacji", "Standard dostawy", "Nadawca", "NIP", "Cennik",
    "Miejsce nadania", "Data nadania", "Data doręczenia", "Odbiorca miasto",
    "Odbiorca kod", "Ilość szt", "Waga", "Objętość", "Wartość towaru",
    "Kwota pobrania", "Typ usługi", "Opłata za usługę", "Cena łączna za zlecenie",
    "Waluta",
]


def write_zadbano_xlsx(path: Path, lines: list[ZLine], summary: dict) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "zlecenia"
    ws.append(ZADBANO_HEADERS)

    lp = 0
    for ln in lines:
        lp += 1
        total = round(sum(c.amount for c in ln.comps), 2)
        for j, comp in enumerate(ln.comps):
            if j == 0:
                row = [
                    lp, ln.waybill, ln.order_ref, "Zlecenie standardowe",
                    ln.status, ln.standard, "MyBed", "1234567890", "Cennik A",
                    "Warszawa", "2026-02-06", "2026-02-07", ln.city, ln.postcode,
                    ln.qty, ln.weight, ln.volume, 2499.0, 0.0,
                    comp.type_raw, comp.amount, total, "PLN",
                ]
            else:
                # wiersz-komponent: pola nagłówkowe puste
                row = [""] * len(ZADBANO_HEADERS)
                row[19] = comp.type_raw
                row[20] = comp.amount
            ws.append(row)

    # arkusz podsumowanie
    ws2 = wb.create_sheet("podsumowanie")
    ws2.append(["Kategoria", "Kwota netto"])
    for cat, amt in summary["categories"].items():
        ws2.append([cat, amt])
    ws2.append(["Total", summary["total"]])

    # arkusz pozostałe (pusty przykładowy)
    ws3 = wb.create_sheet("pozostałe")
    ws3.append(["LP", "Opis", "Kwota"])

    # arkusz info
    ws4 = wb.create_sheet("info")
    ws4.append(["VendorID", "ZADBANO-001"])
    ws4.append(["BillingPeriodID", "2026-02"])
    ws4.append(["GeneratedAt", "2026-03-01T10:00:00"])

    wb.save(path)


def write_zadbano_broken(clean_path: Path, broken_path: Path) -> None:
    """Skopiuj xlsx i uszkodź `xl/styles.xml` (odwzorowanie realnego pliku)."""
    shutil.copy(clean_path, broken_path)
    # przepisz zip, podmieniając styles.xml na treść, której openpyxl nie łyka
    tmp = broken_path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(broken_path, "r") as zin, \
         zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            if item == "xl/styles.xml":
                zout.writestr(item, _BROKEN_STYLES_XML)
            else:
                zout.writestr(item, zin.read(item))
    tmp.replace(broken_path)


_BROKEN_STYLES_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
    '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
    '<fills count="1"><fill><patternFill patternType="???">'
    '<fgColor rgb="NOTACOLOR"/><bogusChild/></patternFill></fill></fills>'
    '<cellXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellXfs>'
    '</styleSheet>'
)


# ---- PDF (SPT / D&M) ------------------------------------------------------- #
# Wspólny render „flat": kolumny na stałych pozycjach X, wiersze wieloliniowe
# (adres w 2 liniach) scalane po pierwszej kolumnie (Lp.).

def _draw_flat_pdf(path: Path, header_title: str, columns: list[tuple[str, float]],
                   rows: list[list[str]], repeat_pages: int = 1,
                   footer: list[str] | None = None) -> None:
    page = landscape(A4)
    width, height = page
    c = canvas.Canvas(str(path), pagesize=page)

    def render_once():
        rows_per_page = 18
        chunks = [rows[i:i + rows_per_page] for i in range(0, len(rows), rows_per_page)] or [[]]
        for ci, chunk in enumerate(chunks):
            c.setFont(_FONT_BOLD, 9)
            c.drawString(30, height - 30, header_title)
            c.setFont(_FONT_BOLD, 7)
            header_y = height - 50
            # nagłówki wielowyrazowe rysujemy PIONOWO (jak druk Symfonia),
            # żeby etykiety sąsiednich kolumn nie zlewały się w poziomie.
            max_words = max(len(label.split()) for label, _ in columns)
            for label, x in columns:
                for wi, word in enumerate(label.split()):
                    c.drawString(x, header_y - wi * 9, word)
            c.setFont(_FONT, 7)
            y = header_y - max_words * 9 - 6
            for row in chunk:
                # row może mieć podwiersz adresu jako '\n'
                maxlines = max(len(str(cell).split("\n")) for cell in row)
                for li in range(maxlines):
                    for (label, x), cell in zip(columns, row):
                        parts = str(cell).split("\n")
                        text = parts[li] if li < len(parts) else ""
                        c.drawString(x, y, text)
                    y -= 10
                y -= 2
            if footer and ci == len(chunks) - 1:
                y -= 6
                c.setFont(_FONT_BOLD, 8)
                for fl in footer:
                    c.drawString(30, y, fl)
                    y -= 12
            c.showPage()

    for _ in range(repeat_pages):
        render_once()
    c.save()


def write_spt_pdf(path: Path, lines: list[FlatLine]) -> None:
    columns = [
        ("Lp.", 30), ("Stan", 58), ("Data rozliczenia", 118), ("ID zam.", 180),
        ("Nr zam.", 222), ("Typ zamówienia", 320), ("Odbiorca", 402),
        ("Adres", 490), ("Produkty", 578), ("Kwota pobrania", 690),
        ("Cena transportu", 748),
    ]
    rows = []
    for i, ln in enumerate(lines, start=1):
        rows.append([
            str(i), ln.status_label, "2026-02-07", f"ID{9000+i}", ln.order_ref,
            ln.order_type, ln.receiver,
            f"{ln.postcode} {ln.city}\n{ln.city}", ln.products,
            pl(0.0), pl(ln.transport_net),
        ])
    total = round(sum(l.transport_net for l in lines), 2)
    footer = [f"PODSUMOWANIE   Suma Cena transportu: {pl(total)}"]
    _draw_flat_pdf(path, f"MYBED LUTY II  Rozliczenie SPT  Netto: {pl(total)}",
                   columns, rows, repeat_pages=1, footer=footer)


def write_dm_pdf(path: Path, lines: list[FlatLine]) -> None:
    columns = [
        ("L.p.", 30), ("Dane odbiorcy", 60), ("Nr. zamówienia", 250),
        ("Opis przesyłki", 330), ("Parametry paczek", 470),
        ("Liczba paczek", 560), ("Pobranie", 620),
        ("Koszt transportu", 680), ("Usługi", 760),
    ]
    rows = []
    for i, ln in enumerate(lines, start=1):
        rows.append([
            str(i), f"{ln.receiver}\n{ln.postcode} {ln.city}", ln.order_ref,
            ln.products, f"{ln.weight} [kg]\n{ln.volume} [m3]", str(ln.packages),
            pl(0.0), pl(ln.transport_net), pl(ln.services_net),
        ])
    total = round(sum(l.transport_net for l in lines), 2)
    footer = [
        f"RAZEM   Koszt transportu netto: {pl(total)}",
        f"Podsumowanie: netto {pl(total)}",
    ]
    title = (f"Rozliczenie o numerze {DM_SETTLEMENT} dla MyBed w walucie PLN, "
             f"stawka VAT 23% 07.02.2026-07.02.2026")
    # treść zdublowana (te same strony x2) -> test dedupu
    _draw_flat_pdf(path, title, columns, rows, repeat_pages=2, footer=footer)


# ---- Faktury zbiorcze (PDF) ------------------------------------------------ #

def write_invoice_pdf(path: Path, *, invoice_no: str, seller: str, nip: str,
                      issue_date: str, period: str, net: float,
                      settlement_no: str | None = None) -> None:
    vat = round(net * 0.23, 2)
    gross = round(net + vat, 2)
    page = A4
    width, height = page
    c = canvas.Canvas(str(path), pagesize=page)
    c.setFont(_FONT_BOLD, 12)
    c.drawString(40, height - 50, f"Faktura VAT {invoice_no}")
    c.setFont(_FONT, 10)
    y = height - 90
    body = [
        f"Sprzedawca: {seller}",
        f"NIP: {nip}",
        f"Data wystawienia: {issue_date}",
        f"Okres rozliczeniowy: {period}",
    ]
    if settlement_no:
        body.append(f"Dotyczy rozliczenia nr: {settlement_no}")
    body += [
        "",
        f"Wartość netto: {pl(net)} PLN",
        f"VAT 23%: {pl(vat)} PLN",
        f"Wartość brutto: {pl(gross)} PLN",
        f"Termin płatności: 2026-03-21",
    ]
    for row in body:
        c.drawString(40, y, row)
        y -= 18
    c.showPage()
    c.save()


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #

def generate_all(out_dir: Path | None = None) -> dict:
    out_dir = Path(out_dir) if out_dir else FIX_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    ERP.clear()

    zad_lines, zad_summary = build_zadbano()
    dm_lines = build_dm()
    spt_lines = build_spt()
    add_extra_erp()

    write_erp_csv(out_dir / "Eksport.csv")

    zad_path = out_dir / "zadbano_spec.xlsx"
    write_zadbano_xlsx(zad_path, zad_lines, zad_summary)
    write_zadbano_broken(zad_path, out_dir / "zadbano_spec_broken.xlsx")

    write_dm_pdf(out_dir / "dm_settlement.pdf", dm_lines)
    write_spt_pdf(out_dir / "spt_spec.pdf", spt_lines)

    write_invoice_pdf(out_dir / "invoice_spt.pdf", invoice_no=SPT_INVOICE,
                      seller="SPT Logistics Sp. z o.o.", nip="1112223344",
                      issue_date="2026-03-01", period=PERIOD, net=SPT_TARGET)
    write_invoice_pdf(out_dir / "invoice_zadbano.pdf", invoice_no=ZADBANO_INVOICE,
                      seller="Zadbano Sp. z o.o.", nip="5556667788",
                      issue_date="2026-03-01", period=PERIOD, net=ZADBANO_TARGET)
    write_invoice_pdf(out_dir / "invoice_dm.pdf", invoice_no=DM_INVOICE,
                      seller="D&M Trans", nip="9998887766",
                      issue_date="2026-02-08", period=PERIOD, net=DM_TARGET,
                      settlement_no=DM_SETTLEMENT)

    manifest = {
        "period": PERIOD,
        "invoices": {
            "SPT": {"invoice_no": SPT_INVOICE, "net": SPT_TARGET},
            "ZADBANO": {"invoice_no": ZADBANO_INVOICE, "net": ZADBANO_TARGET},
            "DM_TRANS": {"invoice_no": DM_INVOICE, "net": DM_TARGET,
                         "settlement_no": DM_SETTLEMENT},
        },
        "golden": {
            "spt_double_core": "46046",
            "spt_double_amounts": [259.60, 583.00],
            "cross_carrier_core": "47559",
            "cross_carrier_amounts": {"ZADBANO": 568.52, "DM_TRANS": 455.00},
            "zadbano_failed_count": 12,
            "zadbano_failed_sum": 5251.37,
            "zadbano_cancelled_count": 3,
        },
        "erp_orders": len(ERP),
        "zadbano_lines": len(zad_lines),
        "dm_lines": len(dm_lines),
        "spt_lines": len(spt_lines),
    }
    (out_dir / "manifest.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False), encoding="utf-8")
    return manifest


if __name__ == "__main__":
    m = generate_all()
    print(json.dumps(m, indent=2, ensure_ascii=False))
