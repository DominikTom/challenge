"""Testy end-to-end pipeline'u i generowania wyjść (§7)."""

import json

from openpyxl import load_workbook

from transport_audit.report import generate_reports, build_enriched_rows
from transport_audit.core.supabase_loader import build_fact_rows


REQUIRED_TABS = {
    "Cross_carrier", "Duble", "Przeplaty", "Nieudane_obciazone",
    "Service_mismatch", "Surcharge_outliers", "Osierocone", "Nieobciazone",
    "Podsumowanie_per_przewoznik", "Walidacja_vs_reczne",
}


def test_generate_reports_creates_files(pipeline_result, tmp_path):
    paths = generate_reports(pipeline_result, tmp_path)
    from pathlib import Path
    for key, p in paths.items():
        assert Path(p).exists(), f"brak pliku {key}: {p}"
    assert (tmp_path / "enriched_orders.csv").exists()
    assert (tmp_path / "audit_report.xlsx").exists()
    assert (tmp_path / "reference_tariff.json").exists()


def test_audit_report_has_required_tabs(pipeline_result, tmp_path):
    generate_reports(pipeline_result, tmp_path)
    wb = load_workbook(tmp_path / "audit_report.xlsx", read_only=True)
    tabs = set(wb.sheetnames)
    missing = REQUIRED_TABS - tabs
    assert not missing, f"brak zakładek: {missing}"
    # Cross_carrier musi być pierwszą zakładką (priorytet WYSOKI)
    assert wb.sheetnames[0] == "Cross_carrier"


def test_enriched_fills_only_empty_transport_invoice(pipeline_result):
    rows = build_enriched_rows(pipeline_result)
    by_num = {r["Numer"]: r for r in rows}
    # 47559 miał ręczny wpis -> zostaje, nie nadpisujemy
    assert by_num["Shoper47559-1"]["Faktura transportowa"] == "307/02/2026/TR ZADBANO"
    assert by_num["Shoper47559-1"]["transport_invoice_validation"] == "OK"
    # zamówienie bez ręcznego wpisu, dopasowane -> uzupełnione auto
    filled = [r for r in rows if r["carrier"] and not r["existing_transport_invoice"]]
    assert filled
    assert all(r["Faktura transportowa"] for r in filled)


def test_enriched_cross_carrier_columns(pipeline_result):
    rows = build_enriched_rows(pipeline_result)
    r = next(r for r in rows if r["Numer"] == "Shoper47559-1")
    assert set(r["carrier"].split(",")) == {"ZADBANO", "DM_TRANS"}
    assert r["real_transport_cost_net"] == 1023.52
    assert "CROSS_CARRIER" in r["audit_flags"]


def test_cost_breakdown_json_valid(pipeline_result):
    rows = build_enriched_rows(pipeline_result)
    r = next(r for r in rows if r["Numer"] == "Shoper47559-1")
    breakdown = json.loads(r["cost_breakdown_json"])
    carriers = {b["carrier"] for b in breakdown}
    assert carriers == {"ZADBANO", "DM_TRANS"}


def test_norm_period_from_docs():
    from transport_audit.core.pipeline import _norm_period
    assert _norm_period("FS/38/02/2026") == "2026-02"
    assert _norm_period("307/02/2026/TR") == "2026-02"
    assert _norm_period("07/02/2026") == "2026-02"
    assert _norm_period("FS/133/06/2026 D&M TRANS") == "2026-06"
    assert _norm_period("2026") is None
    assert _norm_period("") is None


def test_detect_period_from_uploaded_docs(pipeline_result):
    # okres wykryty z faktur/zestawień, bez podawania ręcznie
    from transport_audit.core.pipeline import detect_period
    assert detect_period(pipeline_result.invoices, pipeline_result.deliveries) == "2026-02"


def test_supabase_fact_rows_keyed(pipeline_result):
    rows = build_fact_rows(pipeline_result)
    assert rows
    keys = [(r["order_core"], r["period"], r["carrier"]) for r in rows]
    assert len(keys) == len(set(keys))  # unikalny klucz (order_core, period, carrier)
